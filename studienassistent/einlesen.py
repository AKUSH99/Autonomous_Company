"""Kursunterlagen einlesen: PDF, PowerPoint, Word und Markdown → Abschnitte mit Modul, Datei und Seite.

    python -m studienassistent einlesen <Ordner> [<Ordner> …]

Das Modul ist der erste Ordner unterhalb des angegebenen Ordners (bei Teams-Ablagen: der Team-Ordner). Jede Seite bzw.
Folie wird in Abschnitte von etwa 900 Zeichen zerlegt (mit Überlappung), damit Treffer genau zitiert werden können.
Ergebnis: <daten>/abschnitte.jsonl
"""
from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path

ENDUNGEN = {".pdf", ".pptx", ".docx", ".md", ".txt"}
ABSCHNITT_ZEICHEN = 900
UEBERLAPPUNG = 150


@dataclass
class Abschnitt:
    id: str
    modul: str
    datei: str      # Pfad relativ zum eingelesenen Ordner
    seite: int      # Seite bzw. Folie (1-basiert); 0 bei Text ohne Seiten
    text: str

    def quelle(self) -> str:
        name = Path(self.datei).name
        return f"{self.modul} · {name}" + (f", S. {self.seite}" if self.seite else "")


def seiten_lesen(pfad: Path) -> list[str]:
    """Text je Seite bzw. Folie. Unlesbare Dateien liefern eine leere Liste."""
    try:
        if pfad.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            return [(s.extract_text() or "") for s in PdfReader(str(pfad)).pages]
        if pfad.suffix.lower() == ".pptx":
            from pptx import Presentation
            seiten = []
            for folie in Presentation(str(pfad)).slides:
                teile = [f.text_frame.text for f in folie.shapes if f.has_text_frame]
                if folie.has_notes_slide:
                    teile.append(folie.notes_slide.notes_text_frame.text)
                seiten.append("\n".join(teile))
            return seiten
        if pfad.suffix.lower() == ".docx":
            from docx import Document
            doc = Document(str(pfad))
            teile = [p.text for p in doc.paragraphs]
            for tabelle in doc.tables:
                for zeile in tabelle.rows:
                    teile.append(" | ".join(z.text for z in zeile.cells))
            return ["\n".join(teile)]
        return [pfad.read_text(encoding="utf-8", errors="replace")]
    except Exception:  # noqa: BLE001 – eine kaputte Datei soll das Einlesen nicht stoppen
        return []


def zerlegen(text: str) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) <= ABSCHNITT_ZEICHEN:
        return [text] if text else []
    teile, start = [], 0
    while start < len(text):
        ende = min(len(text), start + ABSCHNITT_ZEICHEN)
        if ende < len(text):  # möglichst an einem Absatz oder Satzende trennen
            umbruch = max(text.rfind("\n", start + ABSCHNITT_ZEICHEN // 2, ende), text.rfind(". ", start + ABSCHNITT_ZEICHEN // 2, ende))
            if umbruch > start:
                ende = umbruch + 1
        teile.append(text[start:ende].strip())
        if ende >= len(text):
            break
        start = max(ende - UEBERLAPPUNG, start + 1)
    return [t for t in teile if t]


def modul_von(relativ: Path) -> str:
    teile = relativ.parts
    if len(teile) > 2 and teile[0].lower() == "teams":
        return teile[1].replace("_", " ")
    return teile[0].replace("_", " ") if len(teile) > 1 else "Allgemein"


def einlesen(ordner: list[Path]) -> list[Abschnitt]:
    abschnitte: list[Abschnitt] = []
    for wurzel in ordner:
        for pfad in sorted(Path(wurzel).rglob("*")):
            if not pfad.is_file() or pfad.suffix.lower() not in ENDUNGEN or pfad.name.endswith(".provenance.json"):
                continue
            relativ = pfad.relative_to(wurzel)
            modul = modul_von(relativ)
            seiten = seiten_lesen(pfad)
            for nr, text in enumerate(seiten, 1):
                for k, teil in enumerate(zerlegen(text)):
                    if len(teil) < 40:  # leere Titelfolien u. ä.
                        continue
                    abschnitte.append(Abschnitt(id=f"{relativ.as_posix()}#{nr}.{k}", modul=modul, datei=relativ.as_posix(),
                                                seite=nr if len(seiten) > 1 or pfad.suffix.lower() == ".pdf" else 0, text=teil))
    return abschnitte


def speichern(abschnitte: list[Abschnitt], ziel: Path) -> None:
    ziel.parent.mkdir(parents=True, exist_ok=True)
    with ziel.open("w", encoding="utf-8") as f:
        for a in abschnitte:
            f.write(json.dumps(asdict(a), ensure_ascii=False) + "\n")


def laden(datei: Path) -> list[Abschnitt]:
    return [Abschnitt(**json.loads(z)) for z in datei.read_text(encoding="utf-8").splitlines() if z.strip()]
